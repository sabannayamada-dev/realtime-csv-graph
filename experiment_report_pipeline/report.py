from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor

from .data_io import read_data


def _set_font(run, name: str, size: float, bold: bool = False) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def _shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_borders(cell, color: str = "D9D9D9") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), color)


def _add_page_number(section) -> None:
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def _prepare_document(settings: dict[str, Any]) -> Document:
    document = Document()
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(float(settings.get("top_margin_mm", 22)))
    section.bottom_margin = Mm(float(settings.get("bottom_margin_mm", 20)))
    section.left_margin = Mm(float(settings.get("left_margin_mm", 25)))
    section.right_margin = Mm(float(settings.get("right_margin_mm", 25)))
    _add_page_number(section)

    font_name = settings.get("font", "Yu Mincho")
    body_size = float(settings.get("body_size", 10.5))
    normal = document.styles["Normal"]
    normal.font.name = font_name
    normal.font.size = Pt(body_size)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    title_style = document.styles["Title"]
    title_style.font.name = font_name
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    title_p_pr = title_style.element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)
    for level, size in ((1, 14), (2, 12), (3, 11)):
        style = document.styles[f"Heading {level}"]
        style.font.name = font_name
        style.font.size = Pt(float(settings.get(f"heading{level}_size", size)))
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    return document


def _add_cover(document: Document, report: dict[str, Any], settings: dict[str, Any]) -> None:
    font_name = settings.get("font", "Yu Mincho")
    paragraph = document.add_paragraph()
    paragraph.style = document.styles["Title"]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(36)
    run = paragraph.add_run(report.get("title", "実験レポート"))
    _set_font(run, font_name, float(settings.get("title_size", 18)), bold=True)

    subtitle = report.get("subtitle")
    if subtitle:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(str(subtitle))
        _set_font(run, font_name, 12)

    metadata = report.get("metadata", [])
    if metadata:
        document.add_paragraph()
        table = document.add_table(rows=0, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for item in metadata:
            cells = table.add_row().cells
            cells[0].text = str(item.get("label", ""))
            cells[1].text = str(item.get("value", ""))
            for cell in cells:
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                _set_cell_borders(cell)
                for cell_paragraph in cell.paragraphs:
                    for cell_run in cell_paragraph.runs:
                        _set_font(cell_run, font_name, 10.5, bold=cell is cells[0])
    document.add_page_break()


def _replace_references(text: str, references: dict[str, str]) -> str:
    pattern = re.compile(r"\[\[(FIG|TABLE):([^\]]+)\]\]")
    return pattern.sub(
        lambda match: references.get(f"{match.group(1)}:{match.group(2)}", match.group(0)),
        text,
    )


def _add_paragraph(document: Document, text: str, settings: dict[str, Any]) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Pt(float(settings.get("body_size", 10.5)))
    paragraph.paragraph_format.line_spacing = float(settings.get("line_spacing", 1.15))
    paragraph.add_run(text)


def _add_figure(document: Document, path: Path, caption: str, number: int) -> None:
    if not path.exists():
        raise FileNotFoundError(f"図のファイルが見つかりません: {path}")
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(path), width=Cm(14.5))
    caption_paragraph = document.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.keep_with_next = True
    caption_run = caption_paragraph.add_run(f"図 {number}  {caption}")
    caption_run.bold = True


def _add_table(document: Document, config: dict[str, Any], base_dir: Path, number: int) -> None:
    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.keep_with_next = True
    caption_run = caption.add_run(f"表 {number}  {config.get('caption', '')}")
    caption_run.bold = True

    frame = read_data((base_dir / config["source"]).resolve(), header=config.get("header", 0))
    columns = config.get("columns") or list(frame.columns)
    missing = [column for column in columns if column not in frame.columns]
    if missing:
        raise ValueError(f"表に指定した列がありません: {', '.join(missing)}")
    frame = frame.loc[:, columns].head(int(config.get("max_rows", 100)))

    table = document.add_table(rows=1, cols=len(columns))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    header_properties = table.rows[0]._tr.get_or_add_trPr()
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    header_properties.append(repeat_header)
    for index, column in enumerate(columns):
        cell = table.rows[0].cells[index]
        cell.text = str(config.get("labels", {}).get(column, column))
        _shade_cell(cell, "D9EAF7")
    for _, row in frame.iterrows():
        cells = table.add_row().cells
        for index, column in enumerate(columns):
            value = row[column]
            cells[index].text = "" if value != value else str(value)
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            _set_cell_borders(cell)
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.size = Pt(9)


def generate_report(
    report: dict[str, Any],
    base_dir: Path,
    output_dir: Path,
    graph_paths: dict[str, Path],
) -> Path:
    settings = report.get("settings", {})
    document = _prepare_document(settings)
    _add_cover(document, report, settings)

    figure_specs: list[dict[str, Any]] = []
    table_specs: list[dict[str, Any]] = []
    for section in report.get("sections", []):
        figure_specs.extend(section.get("figures", []))
        table_specs.extend(section.get("tables", []))
    references: dict[str, str] = {}
    for number, figure in enumerate(figure_specs, start=1):
        if figure.get("id"):
            references[f"FIG:{figure['id']}"] = f"図 {number}"
    for number, table in enumerate(table_specs, start=1):
        if table.get("id"):
            references[f"TABLE:{table['id']}"] = f"表 {number}"

    figure_number = 1
    table_number = 1
    heading_counters = [0, 0, 0, 0]
    for section in report.get("sections", []):
        heading = section.get("heading")
        if heading:
            level = int(section.get("level", 1))
            if level not in (1, 2, 3):
                raise ValueError(f"見出しレベルは1から3で指定してください: {level}")
            heading_counters[level] += 1
            for deeper_level in range(level + 1, 4):
                heading_counters[deeper_level] = 0
            heading_text = str(heading)
            if report.get("number_headings", True):
                number = ".".join(str(value) for value in heading_counters[1 : level + 1])
                heading_text = f"{number}  {heading_text}"
            document.add_heading(heading_text, level=level)
        for paragraph in section.get("paragraphs", []):
            _add_paragraph(document, _replace_references(str(paragraph), references), settings)
        for figure in section.get("figures", []):
            if "graph" in figure:
                try:
                    path = graph_paths[figure["graph"]]
                except KeyError as exc:
                    raise ValueError(f"未生成のグラフIDです: {figure['graph']}") from exc
            else:
                path = (base_dir / figure["path"]).resolve()
            _add_figure(document, path, str(figure.get("caption", "")), figure_number)
            figure_number += 1
        for table in section.get("tables", []):
            _add_table(document, table, base_dir, table_number)
            table_number += 1

    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / report.get("output", "experiment_report.docx")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return output_path.resolve()
