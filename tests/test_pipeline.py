from __future__ import annotations

import json
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from experiment_report_pipeline.cli import run_pipeline


def test_sample_pipeline(tmp_path: Path) -> None:
    root = Path(__file__).resolve().parents[1]
    source = json.loads((root / "examples" / "pipeline_config.json").read_text(encoding="utf-8"))
    source["graphs"][0]["input"] = str((root / "examples" / "sample_data.csv").resolve())
    source["report"]["sections"][1]["tables"][0]["source"] = str(
        (root / "examples" / "sample_data.csv").resolve()
    )
    config_path = tmp_path / "config.json"
    config_path.write_text(json.dumps(source, ensure_ascii=False), encoding="utf-8")

    result = run_pipeline(config_path, tmp_path / "outputs")
    graph_path = Path(result["graphs"]["voltage_current"])
    report_path = Path(result["report"])
    assert graph_path.exists() and graph_path.stat().st_size > 10_000
    assert report_path.exists() and report_path.stat().st_size > 20_000

    document = Document(report_path)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "電圧電流特性実験レポート" in text
    assert "1  目的" in text
    assert "2  実験結果" in text
    assert "3  考察" in text
    assert "図 1" in text
    assert "表 1" in text
    assert "[[FIG:" not in text
    assert len(document.tables) == 2
    assert document.styles["Title"].element.get_or_add_pPr().find(qn("w:pBdr")) is None
    with zipfile.ZipFile(report_path) as archive:
        media = [name for name in archive.namelist() if name.startswith("word/media/")]
    assert media
